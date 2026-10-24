"""ITC test runner with local outcome-first Save-as learning (two runtime files)."""
import argparse
import hashlib
import json
import queue
import sys
import threading
import time
import uuid
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

R=Path(__file__).resolve().parent


def load_config(path):
    from openpyxl import load_workbook
    workbook=load_workbook(path,read_only=True,data_only=True)
    try:
        sheet=workbook['Config'] if 'Config' in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        allowed={'ITC_EXE','ITC_USERNAME','ITC_PASSWORD','ITC_EXPORT_PATH','ITC_SAVE_DELAY_MS','ITC_WORKDIR'}
        cfg={str(row[0]).strip():str(row[1]) if len(row)>1 and row[1] is not None else ''
             for row in sheet.iter_rows(values_only=True) if row and row[0] is not None and str(row[0]).strip() in allowed}
    finally:
        workbook.close()
    if not cfg.get('ITC_EXE'):
        raise ValueError('Config requires ITC_EXE')
    return cfg


def fingerprint(path):
    try:
        digest=hashlib.sha256()
        source=Path(path)
        with source.open('rb') as stream:
            for block in iter(lambda:stream.read(1024*1024),b''):
                digest.update(block)
        return {'path':str(source),'on_disk_sha256':digest.hexdigest(),'size':source.stat().st_size}
    except OSError as error:
        return {'path':str(path),'error':str(error)}


class LocalLog:
    def __init__(self,directory):
        self.directory=directory
        directory.mkdir(parents=True,exist_ok=False)
        self.lock=threading.RLock()
        self.stream=(directory/'events.jsonl').open('x',encoding='utf-8',buffering=1)
        self.failure=None
        self.sequence=0

    def write(self,kind,**fields):
        with self.lock:
            if self.stream.closed:
                return
            self.sequence+=1
            try:
                self.stream.write(json.dumps({'host_sequence':self.sequence,'host_utc':datetime.now(timezone.utc).isoformat(),
                                              'host_monotonic_ns':time.perf_counter_ns(),'kind':kind,**fields},ensure_ascii=False)+'\n')
            except OSError as error:
                self.failure=str(error)

    def save(self,name,value):
        target=self.directory/name
        with target.open('x',encoding='utf-8') as stream:
            json.dump(value,stream,ensure_ascii=False,indent=2)
        return target

    def close(self):
        with self.lock:
            self.stream.close()


HELP='''
x          Existing automation through Filter Enter; wait for report completion yourself.
h / learn  Arm learning, then manually click CHILD File > Save as, observe Select file, Cancel.
j / report Finish learning after Cancel; write discovery report for Copilot.
m / replay Programmatically reopen Select file using learned QSignal; manually Cancel again.
i          Save actual EXE/Qt paths, module inventory and on-disk hashes.
s          Show existing flow status plus compact learning status.
c / l      Existing Connect / Connect+login steps.
a / b      Existing Batch source / invoke Batch.
u / t / k  Existing Batch-to-Filter / typeahead / open Filter.
g / r      Save discovery report without changing learning state.
y / z / e  Disabled in this discovery build; no automatic file export.
q          Dispose and detach; refuses while modal native calls are active.
?          Show this help.
'''


def wait_for(api,predicate,seconds=20):
    deadline=time.monotonic()+seconds
    while time.monotonic()<deadline:
        status=api.status()
        if status.get('error'):
            raise RuntimeError(status['error'])
        if predicate(status):
            return status
        time.sleep(0.05)
    raise TimeoutError('Step did not reach its expected observed state')


def connect_login(sc,cfg,login):
    api=sc.exports_sync
    wait_for(api,lambda st:st.get('action_ready'),30)
    result=api.open()
    print('[OPEN]',result)
    if not result.get('ok'):return
    wait_for(api,lambda st:st.get('ok_ready'))
    result=api.pressok()
    print('[OK]',result)
    if not result.get('ok'):return
    wait_for(api,lambda st:st.get('authorization_seen'))
    if login:
        if not cfg.get('ITC_USERNAME') or not cfg.get('ITC_PASSWORD'):
            raise ValueError('Login requires ITC_USERNAME and ITC_PASSWORD in config.xlsx')
        wait_for(api,lambda st:st.get('authorization_mapping_ready'))
        print('[LOGIN]',api.login(cfg['ITC_USERNAME'],cfg['ITC_PASSWORD']))


def batch_to_filter(sc):
    api=sc.exports_sync
    if not api.status().get('typeahead_table'):
        raise ValueError('Batch table not mapped')
    boot=api.bootstraptypeahead()
    print('[BOOTSTRAP]',boot)
    if not boot.get('ok'):return
    if not boot.get('already_ready'):
        wait_for(api,lambda st:st.get('typeahead_bootstrap_sent') or st.get('typeahead_bootstrap_error'),5)
    print('[ACTIVATE CURRENT]',api.activatecurrentoperation())
    wait_for(api,lambda st:st.get('authorization_list_seen') and st.get('filter_ready'))
    print('[FILTER]',api.invokefilter())


def summary(report):
    return {key:report.get(key) for key in ('version','capabilities','armed','selected','mapping_live','verified','executing','dialog_open','degraded','hook_errors')}


def save_report(api,log,finish=False):
    result=api.finishsavelearning() if finish else {'ok':True,'report':api.savereport()}
    report=result.get('report')
    if report is None:
        print('[LEARNING]',result)
        log.write('learning_result',result=result)
        return result
    path=log.save(f'discovery-{time.time_ns()}.json',report)
    text=['OUTCOME-FIRST SAVE-AS DISCOVERY',json.dumps(summary(report),ensure_ascii=False,indent=2)]
    for sample in report.get('samples',[]):
        text.append(f"sample={sample.get('id')} mode={sample.get('mode')} status={sample.get('status')} owner_matches_report={sample.get('owner_matches_report')} closed={sample.get('closed')}")
        for candidate in sample.get('candidates',[]):
            text.append('  candidate '+json.dumps(candidate,ensure_ascii=False))
    log.save(path.stem+'-compact.json',{'summary':summary(report),'sample_statuses':[
        {key:s.get(key) for key in ('id','mode','status','owner_matches_report','closed','form_observed_ready')}
        for s in report.get('samples',[])]})
    (log.directory/(path.stem+'.txt')).write_text('\n'.join(text),encoding='utf-8')
    print('[LEARNING]',json.dumps(summary(report),ensure_ascii=False,indent=2))
    print('[REPORT SAVED]',path)
    if finish and result.get('ok'):
        print('[NEXT] Use m to reopen Select file; observe it and Cancel. This does not export a file.')
    return result


def save_modules(api,log):
    result=api.moduleaudit()
    result['on_disk_files']=[fingerprint(m['path']) for m in [result.get('main_module'),*result.get('qt_modules',[])] if m]
    path=log.save(f'modules-{time.time_ns()}.json',result)
    print('[MODULES SAVED]',path)
    return result


def replay(api,log):
    queued=api.invokeqsignalsave()
    print('[REPLAY QUEUE]',queued)
    log.write('replay_queue',result=queued)
    if not queued.get('ok'):return
    command_id=queued['command_id']
    deadline=time.monotonic()+20
    while time.monotonic()<deadline:
        reply=api.savecommandstatus(command_id)
        command=reply.get('command',{})
        if command.get('outcome_seen'):
            print('[REPLAY OUTCOME] Select file/owner/form observed. Check the dialog and Cancel manually; then j.')
            log.write('replay_outcome_observed',command=command)
            return
        if command.get('finished'):
            print('[REPLAY RESULT]',json.dumps(command,ensure_ascii=False,indent=2))
            return
        time.sleep(0.05)
    print('[REPLAY TIMEOUT] No success declared. Use s/j; do not retry until the pending native call is resolved.')
    log.write('replay_host_timeout',command_id=command_id)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config',type=Path,default=R/'config.xlsx')
    parser.add_argument('--cwd-exe',action='store_true',help='Controlled experiment: use EXE folder as child working directory')
    args=parser.parse_args()
    import frida
    cfg=load_config(args.config)
    exe=Path(cfg['ITC_EXE']).expanduser().resolve()
    if not exe.is_file():raise ValueError('ITC_EXE does not exist')
    workdir=str(exe.parent) if args.cwd_exe else cfg.get('ITC_WORKDIR') or None
    run_name=datetime.now().strftime('%Y%m%d-%H%M%S')+'-'+uuid.uuid4().hex[:8]
    log=LocalLog(R/'captures'/run_name)
    print('[CAPTURE DIRECTORY]',log.directory)
    log.write('host_context',configured_exe=cfg['ITC_EXE'],resolved_exe=str(exe),python_cwd=str(Path.cwd()),child_cwd_requested=workdir,
              agent=fingerprint(R/'agent.js'),runner=fingerprint(Path(__file__)))
    device=session=script=None
    pid=None
    resumed=False
    detached=threading.Event()
    crashed=threading.Event()
    agent_errors=threading.Event()
    command_queue=queue.Queue()
    exit_code=0

    def on_message(message,_data):
        # The agent emits pointer/connection metadata, not QLineEdit values.
        log.write('frida',message=message)
        if message.get('type')=='error':
            agent_errors.set()
            print('[FRIDA ERROR]',message.get('stack',message),flush=True)
        elif message.get('type')=='send':
            payload=message.get('payload',{})
            if payload.get('type')=='log':
                print('[flow]',payload.get('message'),flush=True)
            else:
                event=payload.get('event',{})
                if event.get('kind','').startswith('sl_') or event.get('kind','').endswith('error'):
                    print('[EVT]',json.dumps(event,ensure_ascii=False),flush=True)
                else:
                    print('[EVT]',event.get('kind'),'thread=',event.get('thread_id'),flush=True)

    def on_detached(reason,crash=None):
        log.write('session_detached',reason=reason,crash=str(crash) if crash else None)
        if crash:crashed.set()
        detached.set()

    def input_loop():
        while True:
            try:value=input('> ').strip().lower()
            except EOFError:
                command_queue.put('eof');return
            command_queue.put(value)
            # Remain available if q is refused while a modal call is active.

    try:
        device=frida.get_local_device()
        kwargs={'cwd':workdir} if workdir else {}
        pid=device.spawn([str(exe)],**kwargs)
        print('[SPAWNED PID]',pid)
        log.write('spawned',pid=pid)
        session=device.attach(pid)
        session.on('detached',on_detached)
        script=session.create_script((R/'agent.js').read_text('utf-8'))
        script.on('message',on_message)
        script.load()
        device.resume(pid)
        resumed=True
        api=script.exports_sync
        wait_for(api,lambda st:st.get('initialized'),30)
        save_modules(api,log)
        print(HELP)
        threading.Thread(target=input_loop,daemon=True).start()
        while not detached.is_set():
            if log.failure:raise OSError('Cannot write diagnostic log: '+log.failure)
            try:command=command_queue.get(timeout=0.2)
            except queue.Empty:continue
            log.write('operator_command',command=command)
            try:
                if command in ('?','help'):print(HELP)
                elif command=='x':
                    if not cfg.get('ITC_USERNAME') or not cfg.get('ITC_PASSWORD'):raise ValueError('x requires login config')
                    run_to_filter(script,cfg)
                elif command in ('c','l'):connect_login(script,cfg,command=='l')
                elif command=='a':print(api.armbatch())
                elif command=='b':print(api.invokebatch())
                elif command=='u':batch_to_filter(script)
                elif command=='t':print(api.invoketypeahead('on-h'))
                elif command=='k':print(api.invokefilter())
                elif command in ('h','learn'):print('[ARM]',api.learnsave())
                elif command in ('j','report'):save_report(api,log,finish=True)
                elif command in ('g','r'):save_report(api,log)
                elif command in ('m','replay'):replay(api,log)
                elif command=='i':save_modules(api,log)
                elif command=='s':print(json.dumps({'flow':api.status(),'learning':summary(api.savereport())},ensure_ascii=False,indent=2))
                elif command in ('y','z','e','v','p','w'):print('[DISABLED] This build discovers/replays opening only. Use x then h/j/m.')
                elif command in ('q','eof'):
                    save_report(api,log)
                    result=api.safedispose()
                    ack_deadline=time.monotonic()+2
                    while result.get('error')=='WAITING_GUI_DISPOSE_ACK' and time.monotonic()<ack_deadline:
                        time.sleep(0.05)
                        result=api.safedispose()
                    print('[DISPOSE]',result)
                    if result.get('ok'):
                        session.detach();break
                    print('[WAIT] Close modal dialogs or generate a normal GUI event in ITC, then q again. Do not start another automation command after shutdown has begun.')
                    if command=='eof':
                        raise RuntimeError('EOF while native calls remain active; keep runner alive and close the target dialog')
                elif command:print('Unknown command. Use ?')
            except (ValueError,TimeoutError) as error:
                print('[STEP ERROR]',error)
                log.write('step_error',command=command,error=str(error))
    except KeyboardInterrupt:
        print('[INTERRUPTED] Attempting safe dispose.')
    except Exception as error:
        exit_code=2
        print('[RUN ERROR]',error,file=sys.stderr)
        log.write('runner_error',error=str(error))
    finally:
        cfg.pop('ITC_PASSWORD',None)
        if pid is not None and not resumed:
            try:device.kill(pid)
            except Exception as error:log.write('suspended_cleanup_error',error=str(error));exit_code=2
        if session is not None and not detached.is_set():
            # Keep the agent alive until its native timer/call has become idle.
            while not detached.is_set():
                try:
                    result=script.exports_sync.safedispose() if script is not None else {'ok':True}
                    if result.get('ok'):
                        session.detach();break
                    print('[CLEANUP WAIT]',result,'- close target modal dialogs; do not force-unload this agent.',flush=True)
                    detached.wait(2)
                except KeyboardInterrupt:
                    print('[CLEANUP WAIT] Agent retained until target is idle or process exits.',flush=True)
                except Exception as error:
                    log.write('cleanup_rpc_error',error=str(error))
                    print('[CLEANUP WAIT] Cannot confirm disposal; close the owned ITC process to finish.',flush=True)
                    detached.wait(2)
        if crashed.is_set() or agent_errors.is_set():exit_code=2
        log.write('runner_finished',exit_code=exit_code)
        log.close()
    print('[LOCAL LOG]',log.directory)
    return exit_code


# run_to_filter is embedded below from the supplied working runner.

def run_to_filter(sc, cfg):
    username=cfg['ITC_USERNAME']; password=cfg['ITC_PASSWORD']
    print('[FULL] Step 1/6: open Connection select')
    end=time.time()+30
    while time.time()<end and not sc.exports_sync.status().get('action_ready'):time.sleep(.05)
    opened=sc.exports_sync.open();print('[OPEN]',opened)
    if not opened.get('ok'):return False
    end=time.time()+20
    while time.time()<end and not sc.exports_sync.status().get('ok_ready'):time.sleep(.05)
    if not sc.exports_sync.status().get('ok_ready'):
     print('[FULL ERROR] Connection OK not mapped',json.dumps(sc.exports_sync.status(),indent=2));return False
    print('[FULL] Step 2/6: open Authorization')
    pressed=sc.exports_sync.pressok();print('[OK]',pressed)
    if not pressed.get('ok'):return False
    end=time.time()+20
    while time.time()<end and not sc.exports_sync.status().get('authorization_mapping_ready'):time.sleep(.05)
    if not sc.exports_sync.status().get('authorization_mapping_ready'):
     print('[FULL ERROR] Authorization controls not mapped',json.dumps(sc.exports_sync.status(),indent=2));return False
    print('[FULL] Step 3/6: login from config.xlsx')
    logged=sc.exports_sync.login(username,password);print('[LOGIN]',logged)
    if not logged.get('ok'):return False
    end=time.time()+30
    while time.time()<end:
     st=sc.exports_sync.status()
     # Batch tooltip is populated after successful login/main UI construction.
     if st.get('gui_thread') and not st.get('pending_login'):break
     time.sleep(.05)
    time.sleep(2)
    print('[FULL] Step 4/6: discover and open Batch operations')
    armed=sc.exports_sync.armbatch();print('[BATCH ARM]',armed)
    if not armed.get('cached_tooltip_widget'):
     print('[FULL ERROR] Batch operations tooltip source not discovered',json.dumps(sc.exports_sync.status(),indent=2));return False
    batch=sc.exports_sync.invokebatch();print('[BATCH OPEN]',batch)
    if not batch.get('ok'):return False
    end=time.time()+15
    while time.time()<end and not sc.exports_sync.status().get('typeahead_table'):time.sleep(.05)
    st=sc.exports_sync.status()
    if not st.get('typeahead_table'):
     print('[FULL ERROR] Batch table not mapped',json.dumps(st,indent=2));return False
    print('[FULL] Step 5/6: Qt-native type-ahead and open authorization list')
    boot=sc.exports_sync.bootstraptypeahead();print('[BOOTSTRAP]',json.dumps(boot,indent=2))
    if not boot.get('ok'):return False
    end=time.time()+5
    while time.time()<end:
     st=sc.exports_sync.status()
     if st.get('typeahead_bootstrap_sent') or st.get('typeahead_bootstrap_error'):break
     time.sleep(.05)
    st=sc.exports_sync.status()
    if not st.get('typeahead_bootstrap_sent'):
     print('[FULL ERROR] Type-ahead bootstrap not delivered',json.dumps(st,indent=2));return False
    current=sc.exports_sync.activatecurrentoperation();print('[ACTIVATE CURRENT]',current)
    if not current.get('ok'):return False
    end=time.time()+15
    while time.time()<end:
     st=sc.exports_sync.status()
     if st.get('authorization_list_seen') and st.get('filter_ready'):break
     time.sleep(.05)
    st=sc.exports_sync.status()
    if not st.get('authorization_list_seen') or not st.get('filter_ready'):
     print('[FULL ERROR] Authorization list or Filter mapping not ready',json.dumps(st,indent=2));return False
    print('[FULL] Step 6/6: click Filter')
    filt=sc.exports_sync.invokefilter();print('[FILTER]',filt)
    if not filt.get('ok'):return False
    end=time.time()+10
    while time.time()<end and not sc.exports_sync.status().get('filter_form_ready'):time.sleep(.05)
    st=sc.exports_sync.status()
    if not st.get('filter_form_ready'):
     print('[FULL ERROR] Filter form controls not mapped',json.dumps(st,indent=2));return False
    today=date.today();from_value=(today-timedelta(days=30)).strftime('%d%m%Y');to_value=(today-timedelta(days=1)).strftime('%d%m%Y')
    print(f'[FULL] Filter dates: T-30={from_value}, T-1={to_value}')
    submitted=sc.exports_sync.submitfilter(from_value,to_value);print('[FILTER SUBMIT]',submitted)
    if not submitted.get('ok'):return False
    time.sleep(1)
    print('[PRE_SAVE] Filter submitted. Wait for report completion, then h. This is not report-completion acknowledgement.')
    return True

if __name__=='__main__':
    raise SystemExit(main())
